import ipaddress
import re
import socket
import zipfile

from io import BytesIO
from pathlib import Path
from urllib.parse import urljoin, urlsplit, urlunsplit

import pymupdf
import httpx
from bs4 import BeautifulSoup
from docx import Document


MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_HTML_SIZE = 2 * 1024 * 1024
MAX_JOB_TEXT_LENGTH = 50000
MAX_REDIRECTS = 3

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
REDIRECT_CODES = {301, 302, 303, 307, 308}


def _clean_text(text: str) -> str:
    """Normalize extracted text and remove blank-only lines."""
    lines = []

    for line in text.replace("\x00", "").splitlines():
        line = re.sub(r"[ \t]+", " ", line).strip()
        if line:
            lines.append(line)

    return "\n".join(lines).strip()


def extract_job_text_from_file(
    filename: str,
    content: bytes,
) -> str:
    """Extract text from PDF, DOCX, or TXT without saving the file."""

    safe_filename = (
        (filename or "")
        .replace("\\", "/")
        .rsplit("/", 1)[-1]
    )
    extension = Path(safe_filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Unsupported file type. Upload a PDF, DOCX, or TXT."
        )

    if not content:
        raise ValueError("The uploaded file is empty.")

    if len(content) > MAX_FILE_SIZE:
        raise ValueError("Maximum upload size is 5 MB.")

    try:
        if extension == ".pdf":
            if not content.startswith(b"%PDF-"):
                raise ValueError("The file is not a valid PDF.")

            with pymupdf.open(stream=content, filetype="pdf") as document:
                if document.needs_pass:
                    raise ValueError("Password-protected PDFs are not supported.")

                if len(document) > 100:
                    raise ValueError("PDFs cannot contain more than 100 pages.")

                extracted = "\n".join(page.get_text() for page in document)

        elif extension == ".docx":
            buffer = BytesIO(content)

            if not zipfile.is_zipfile(buffer):
                raise ValueError("The file is not a valid DOCX.")

            buffer.seek(0)

            with zipfile.ZipFile(buffer) as archive:
                names = archive.namelist()

                if "word/document.xml" not in names:
                    raise ValueError("The file is not a valid DOCX.")

                # Guard against unusually large compressed documents.
                expanded_size = sum(
                    item.file_size for item in archive.infolist()
                )
                if expanded_size > 20 * 1024 * 1024:
                    raise ValueError(
                        "The expanded DOCX document is too large."
                    )

            document = Document(BytesIO(content))
            parts = [
                paragraph.text
                for paragraph in document.paragraphs
            ]

            for table in document.tables:
                for row in table.rows:
                    parts.append(
                        " | ".join(cell.text for cell in row.cells)
                    )

            extracted = "\n".join(parts)

        else:
            try:
                extracted = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                extracted = content.decode(
                    "cp1252",
                    errors="replace",
                )

    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            "Could not read this document. Check that it is valid."
        ) from exc

    text = _clean_text(extracted)

    if len(text) < 20:
        raise ValueError(
            "No readable job description was found in the file. "
            "Scanned PDFs requiring OCR are not supported yet."
        )

    if len(text) > MAX_JOB_TEXT_LENGTH:
        raise ValueError(
            "Extracted text exceeds 50000 characters. "
            "Upload a shorter job description."
        )

    return text


def _validate_public_https_url(url: str) -> str:
    """
    Accept public HTTPS hostnames only.

    Blocks credentials, IP-literal URLs, nonstandard ports, local
    hostnames, and hostnames resolving to non-public IP addresses.
    """
    try:
        parsed = urlsplit(url.strip())
        hostname = parsed.hostname
        port = parsed.port
    except ValueError as exc:
        raise ValueError("Invalid job URL.") from exc

    if parsed.scheme.lower() != "https":
        raise ValueError("Only HTTPS job URLs are supported.")

    if not hostname or parsed.username or parsed.password:
        raise ValueError("Invalid job URL.")

    if port not in (None, 443):
        raise ValueError("Only the standard HTTPS port is allowed.")

    try:
        hostname = hostname.encode("idna").decode("ascii").lower()
    except UnicodeError as exc:
        raise ValueError("Invalid URL hostname.") from exc

    hostname = hostname.rstrip(".")

    if (
        "." not in hostname
        or hostname in {"localhost"}
        or hostname.endswith((".local", ".internal", ".test"))
    ):
        raise ValueError("Local or private hostnames are not allowed.")

    # Do not accept IP-literal URLs.
    try:
        ipaddress.ip_address(hostname)
    except ValueError:
        pass
    else:
        raise ValueError("Use a public HTTPS domain, not an IP address.")

    try:
        addresses = socket.getaddrinfo(
            hostname,
            443,
            type=socket.SOCK_STREAM,
        )
    except OSError as exc:
        raise ValueError(
            "The job website hostname could not be resolved."
        ) from exc

    if not addresses:
        raise ValueError("The job website hostname could not be resolved.")

    for address in addresses:
        try:
            ip = ipaddress.ip_address(address[4][0].split("%")[0])
        except ValueError as exc:
            raise ValueError("Invalid destination address.") from exc

        if not ip.is_global:
            raise ValueError(
                "URLs resolving to private or reserved addresses are blocked."
            )

    # Drop fragments and normalize the URL before fetching it.
    return urlunsplit((
        "https",
        hostname,
        parsed.path or "/",
        parsed.query,
        "",
    ))


def extract_job_text_from_url(url: str) -> tuple[str, str, str]:
    """
    Fetch a public HTTPS job page and extract its readable text.

    Returns: (description_text, page_title, final_url)
    """
    current_url = url.strip()

    timeout = httpx.Timeout(10.0, connect=3.0)

    try:
        with httpx.Client(
            follow_redirects=False,
            timeout=timeout,
            trust_env=False,
            headers={
                "User-Agent": "AIResumeCreator/1.0",
                "Accept": "text/html,application/xhtml+xml",
            },
        ) as client:
            html_content = None

            for redirect_number in range(MAX_REDIRECTS + 1):
                current_url = _validate_public_https_url(current_url)

                with client.stream("GET", current_url) as response:
                    if response.status_code in REDIRECT_CODES:
                        if redirect_number >= MAX_REDIRECTS:
                            raise ValueError(
                                "The job page has too many redirects."
                            )

                        location = response.headers.get("location")
                        if not location:
                            raise ValueError(
                                "The job page returned an invalid redirect."
                            )

                        current_url = urljoin(current_url, location)
                        continue

                    if response.status_code != 200:
                        raise ValueError(
                            "The job page could not be retrieved. "
                            "The website may block automated requests."
                        )

                    content_type = response.headers.get(
                        "content-type", ""
                    ).lower()

                    if not (
                        "text/html" in content_type
                        or "application/xhtml+xml" in content_type
                    ):
                        raise ValueError(
                            "The URL must point to an HTML job page."
                        )

                    chunks = []
                    total_size = 0

                    for chunk in response.iter_bytes():
                        total_size += len(chunk)

                        if total_size > MAX_HTML_SIZE:
                            raise ValueError(
                                "The job page is too large to process."
                            )

                        chunks.append(chunk)

                    html_content = b"".join(chunks)
                    break

            if html_content is None:
                raise ValueError("Could not retrieve the job page.")

    except httpx.TimeoutException as exc:
        raise ValueError("The job website timed out.") from exc
    except httpx.RequestError as exc:
        raise ValueError(
            "Could not connect to the job website."
        ) from exc

    html = html_content.decode("utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")

    page_title = (
        soup.title.get_text(" ", strip=True)
        if soup.title
        else ""
    )[:255]

    for element in soup.select(
        "script, style, noscript, svg, nav, footer, form, "
        "button, iframe, template"
    ):
        element.decompose()

    for element in soup.select(
        "[class*='cookie'], [id*='cookie'], "
        "[class*='privacy-banner'], [class*='advertisement']"
    ):
        element.decompose()

    # Prefer the actual job content over the entire page layout.
    candidates = soup.select(
        "main, article, "
        "[class*='job-description'], "
        "[class*='jobDescription'], "
        "[id*='job-description'], "
        "[data-testid*='job-description']"
    )

    candidate_texts = [
        _clean_text(element.get_text("\n", strip=True))
        for element in candidates
    ]
    candidate_texts = [
        text for text in candidate_texts if text
    ]

    if candidate_texts:
        extracted = max(candidate_texts, key=len)
    elif soup.body:
        extracted = _clean_text(
            soup.body.get_text("\n", strip=True)
        )
    else:
        extracted = _clean_text(soup.get_text("\n", strip=True))

    if len(extracted) < 20:
        raise ValueError(
            "No readable job description was found on the page. "
            "Try pasting the job description text instead."
        )

    # Job pages can include substantial surrounding site content.
    extracted = extracted[:MAX_JOB_TEXT_LENGTH]

    return extracted, page_title, current_url