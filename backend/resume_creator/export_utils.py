
from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
)


def content_to_text(value):
    """Convert JSON section content into readable text."""
    if value is None:
        return ""

    if isinstance(value, str):
        return value.strip()

    if isinstance(value, (int, float, bool)):
        return str(value)

    if isinstance(value, list):
        lines = []
        for item in value:
            text = content_to_text(item)
            if text:
                lines.append(text)
        return "\n".join(lines)

    if isinstance(value, dict):
        lines = []
        for key, item in value.items():
            text = content_to_text(item)
            if text:
                label = str(key).replace("_", " ").title()
                lines.append(f"{label}: {text}")
        return "\n".join(lines)

    return str(value)


def get_section_content(section):
    return content_to_text(section.content)


def get_section_lines(section):
    """Convert section JSON into readable lines for export."""
    content = section.content

    if content is None:
        return []

    if isinstance(content, list):
        return [
            content_to_text(item).strip()
            for item in content
            if content_to_text(item).strip()
        ]

    if isinstance(content, dict):
        lines = []

        for key, value in content.items():
            if value is None or value == "":
                continue

            label = str(key).replace("_", " ").title()

            if isinstance(value, list):
                if not value:
                    continue

                # Simple skill lists become individual bullet points.
                if all(not isinstance(item, dict) for item in value):
                    lines.append(f"{label}:")
                    lines.extend(
                        f"• {content_to_text(item)}"
                        for item in value
                        if content_to_text(item).strip()
                    )
                else:
                    # Keep nested experience/project objects readable.
                    for item in value:
                        text = content_to_text(item).strip()
                        if text:
                            lines.append(f"{label}: {text}")
            else:
                text = content_to_text(value).strip()
                if text:
                    lines.append(f"{label}: {text}")

        return lines

    text = content_to_text(content).strip()
    return text.splitlines() if text else []


def build_pdf(resume, version):
    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=resume.title,
        author="AI Resume Creator",
    )

    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="ResumeName",
            parent=styles["Title"],
            alignment=TA_CENTER,
            fontSize=19,
            leading=23,
            textColor=colors.HexColor("#17365D"),
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ResumeSection",
            parent=styles["Heading2"],
            fontSize=12,
            leading=15,
            spaceBefore=10,
            spaceAfter=5,
            keepWithNext=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ResumeBody",
            parent=styles["BodyText"],
            fontSize=9.5,
            leading=13,
            spaceAfter=4,
        )
    )

    story = [
        Paragraph(escape(resume.title), styles["ResumeName"]),
        Paragraph(
            f"Version {version.version_number}",
            styles["Normal"],
        ),
        Spacer(1, 8),
    ]

    for section in version.sections.filter(
        is_visible=True
    ).order_by("display_order", "created_at"):
        lines = get_section_lines(section)
        if not lines:
            continue

        section_blocks = [
            Paragraph(escape(section.title), styles["ResumeSection"])
        ]

        for line in lines:
            line = line.strip()

            if not line:
                section_blocks.append(Spacer(1, 3))
            elif line.endswith(":"):
                section_blocks.append(
                    Paragraph(escape(line), styles["ResumeSection"])
                )
            elif line.startswith("• "):
                section_blocks.append(
                    Paragraph(
                        escape(line),
                        styles["ResumeBody"],
                    )
                )
            else:
                section_blocks.append(
                    Paragraph(
                        escape(line),
                        styles["ResumeBody"],
                    )
                )

        story.extend(section_blocks)

    document.build(story)
    buffer.seek(0)
    return buffer


def build_docx(resume, version):
    document = Document()

    section_settings = document.sections[0]
    section_settings.top_margin = Inches(0.65)
    section_settings.bottom_margin = Inches(0.65)
    section_settings.left_margin = Inches(0.75)
    section_settings.right_margin = Inches(0.75)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10)

    title = document.add_heading(resume.title, level=0)
    title.alignment = 1

    document.add_paragraph(f"Version {version.version_number}")

    for section in version.sections.filter(
        is_visible=True
    ).order_by("display_order", "created_at"):
        lines = get_section_lines(section)
        if not lines:
            continue

        document.add_heading(section.title, level=1)

        for line in lines:
            line = line.strip()

            if not line:
                continue

            if line.endswith(":"):
                paragraph = document.add_paragraph()
                run = paragraph.add_run(line)
                run.bold = True
            elif line.startswith("• "):
                paragraph = document.add_paragraph(
                    line[2:],
                    style="List Bullet",
                )
            else:
                document.add_paragraph(line)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
